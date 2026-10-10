"""Builds the BUS5001 Assessment 3 report (APA 7 layout) as a .docx file."""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent / "BUS5001_A3_Rishabh_Pandya_22700778.docx"
Q1S = ROOT / "q1_dialogflow" / "screenshots"
Q4S = ROOT / "q4_notebooklm" / "screenshots"

NAME = "Rishabh Jayesh Pandya"
SID = "22700778"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.54)
sec.top_margin = sec.bottom_margin = Cm(2.54)

st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.3

fig_n = 0
tab_n = 0
REGFILE = Path(__file__).resolve().parent / "reg.json"
try:
    REG = json.load(open(REGFILE))
except Exception:
    REG = {}
NEWREG = {}


def _res(text):
    return re.sub(r"\{(fig|tab):(\w+)\}", lambda m: str(REG.get(m.group(1) + ":" + m.group(2), "?")), text)


def _runs(p, text, size=None, font=None, italic=False):
    text = _res(text)
    parts = re.split(r"(\*\*.+?\*\*|\*.+?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*"):
            r = p.add_run(part[1:-1])
            r.italic = True
        else:
            r = p.add_run(part)
        if italic:
            r.italic = True
        if size:
            r.font.size = Pt(size)
        if font:
            r.font.name = font
            r._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def P(text, align=None, indent=True, after=6, size=None, italic=False, keep=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.27)
    p.paragraph_format.space_after = Pt(after)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if keep:
        p.paragraph_format.keep_with_next = True
    _runs(p, text, size=size, italic=italic)
    return p


def H(level, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if level == 3:
        r.italic = True
    return p


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _shade(cell, hex_fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _borders(table):
    tbl = table._element
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)


def caption(kind, title, key=None):
    global fig_n, tab_n
    if kind == "Table":
        tab_n += 1
        n = tab_n
    else:
        fig_n += 1
        n = fig_n
    if key:
        NEWREG[("tab:" if kind == "Table" else "fig:") + key] = n
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{kind} {n}")
    r.bold = True
    p2 = doc.add_paragraph()
    p2.paragraph_format.keep_with_next = True
    p2.paragraph_format.space_after = Pt(4)
    r2 = p2.add_run(title)
    r2.italic = True
    return n


def T(title, header, rows, widths, size=9.5, note=None, numbered=True, key=None):
    if numbered:
        caption("Table", title, key)
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.width = Cm(widths[i])
        c.text = ""
        _shade(c, "D9D9D9")
        _runs(c.paragraphs[0], f"**{h}**", size=size)
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        c.paragraphs[0].paragraph_format.line_spacing = 1.0
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].width = Cm(widths[i])
            cells[i].text = ""
            _runs(cells[i].paragraphs[0], str(v), size=size)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
            cells[i].paragraphs[0].paragraph_format.line_spacing = 1.0
    if note:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        _runs(p, f"*Note.* {note}", size=9.5)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def FIG(title, path, width_cm=14.5, note=None, key=None):
    caption("Figure", title, key)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if note:
        q = doc.add_paragraph()
        _runs(q, f"*Note.* {note}", size=9.5)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def CODE(text, size=8):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = "Courier New"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    r.font.size = Pt(size)


# --------------------------------------------------------------------------------------
# Title page
# --------------------------------------------------------------------------------------
for _ in range(6):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Robotic Process Automation, AI and Cloud-Enabled ESG Analytics")
r.bold = True
r.font.size = Pt(15)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Assessment 3 Written Report")
r.bold = True
for _ in range(2):
    doc.add_paragraph()
for line in [NAME, f"Student number {SID}", "BUS5001", "October 2026"]:
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.add_run(line)
doc.add_paragraph()

T(
    "Submission links",
    ["Item", "Link"],
    [
        ["Q1 demo video (Microsoft Stream)", "[Insert the Stream link after recording and uploading the video]"],
        ["GitHub repository (Q3 and Q4 logs, diagrams, Dialogflow agent)", "[Insert the repository link]"],
    ],
    [6.5, 9.4],
    numbered=False,
)
page_break()

# --------------------------------------------------------------------------------------
# Q1
# --------------------------------------------------------------------------------------
H(1, "Q1 ESG Chatbot in Dialogflow CX")

H(2, "Q1(a) Five ESG chatbot functionalities")
P(
    "The prototype is the ESG Sustainability Support Assistant. It lets staff report sustainability and safety "
    "problems on a campus or office site through a short conversation, and it hands the conversation to a person when "
    "the problem is urgent or the bot cannot help. It was built in Dialogflow CX and tested in the console simulator. "
    "Table {tab:func} lists five functionalities that a chatbot can fulfil, what each does and why it matters to the organisation."
)
T(
    "ESG chatbot functionalities",
    ["#", "Functionality", "ESG pillar", "What the chatbot does", "Why it matters to the organisation", "Status"],
    [
        ["1", "Environmental incident reporting", "Environmental", "Collects the issue type (water leak, energy waste, waste contamination, unsafe disposal), the building and the urgency in three short questions.", "Gives the organisation structured and timely data on leaks, energy waste and contamination. This cuts resource loss and feeds ESG metrics.", "Built and tested"],
        ["2", "Urgent hazard escalation", "Environmental and social", "Recognises danger wording or an answer of immediate risk and passes the user to a person, with a prompt to call 000 first.", "Shortens the time to a response when people or the environment are at risk and supports the organisation’s duty of care.", "Built and tested"],
        ["3", "Accessibility barrier referral", "Social", "Sends blocked ramps, lifts and entrances to the Accessibility and Inclusion Office instead of the general facilities queue.", "Puts the barrier with the team that owns inclusion and creates evidence for social-impact reporting.", "Built and tested"],
        ["4", "Report status guidance", "Governance", "Explains how to follow up a report. The prototype has no ticket system, so it says so and gives a helpdesk address.", "Builds trust in the reporting channel because reporters learn what happens next.", "Built as a stub"],
        ["5", "Privacy and data-handling answers", "Governance", "Explains what the bot collects and points to a confidential channel for conduct concerns.", "Supports transparency and consent, and gives a safe route for sensitive matters.", "Built and tested"],
    ],
    [0.7, 2.6, 1.9, 4.5, 4.4, 1.8],
    size=9,
    key="func",
)

H(2, "Q1(b) Ethics, privacy and governance")
P(
    "An incident bot collects free text from people, and free text can contain names, health details or allegations about "
    "colleagues. The design therefore starts from data minimisation. The bot asks only for the issue type, the building "
    "and the urgency. When a user asks what is collected, the ask.privacy intent answers in plain language and tells the "
    "user not to share personal details. The Australian Privacy Principles expect an organisation to collect only the "
    "personal information it needs (APP 3), to tell people what it collects (APP 5) and to protect it (APP 11) "
    "(*Privacy Act 1988*, Sch 1). The prototype follows these principles in design, but a production release would need a "
    "formal privacy impact assessment."
)
P(
    "Human oversight is the second control. The bot states at the start that it is a virtual assistant and that the user can say "
    "‘human’ at any time. Three failed attempts to understand the user send the conversation to a person, and an "
    "urgent report goes to a person immediately. This keeps the bot from becoming the only route to a response, which matters "
    "for safety reports and for conduct concerns. The bot also stays inside clear boundaries. It uses intent matching and fixed messages written by "
    "the team, so it cannot invent policy, and it does not decide outcomes about people. Table {tab:risk} sets out the main risks and the controls."
)
T(
    "Ethics, privacy and governance risks and controls",
    ["Risk", "Control built into the prototype", "Control needed in production"],
    [
        ["Over-collection of personal data", "Three structured questions only. Privacy intent tells users not to share personal details. Cloud logging and conversation history are switched off in the agent settings.", "Privacy impact assessment. Data retention schedule. Redaction of any personal data found in free text before storage."],
        ["Consent and notice", "Greeting states that the assistant is virtual. Privacy intent explains what is collected and why.", "Collection notice shown before the first message and linked from the chat window, with a clear way to decline and use another channel."],
        ["Lack of transparency and auditability", "Single flow with named routes, so every outcome can be traced to a page and a rule.", "Conversation and routing logs kept with a retention rule, an audit trail of who changed intents or messages, and a record of the purpose of each data item."],
        ["Harm from missed or slow escalation", "Urgent route to a person. Prompt to call 000 first. Fallback to a person after three failed turns.", "Service-level targets for the on-call team. Monitoring of the handoff queue. Regular review of fallback transcripts."],
        ["Bias and unfair treatment", "Urgency depends on the user’s own answer and on the issue type. It does not use the user’s identity or role.", "Testing with varied phrasing, including non-standard English. Review of outcomes by site and by team."],
        ["Retaliation against reporters", "Conduct concerns point to a confidential hotline instead of the general queue.", "Role-based access to transcripts. Whistleblower policy applied to the channel. Audit log of who reads reports."],
        ["Responsible AI boundaries", "No generated free-text answers, no advice on legal or medical matters, and a person is always reachable.", "Written scope statement for the bot. Approval step before any generative feature is added."],
    ],
    [3.2, 6.4, 6.3],
    key="risk",
)

H(2, "Q1(c) Chatbot flow design and implementation")
H(3, "Design thinking and user journey")
P(
    "The flow implements functionality 1, environmental incident reporting, with functionalities 2 and 3 as branches and functionality 5 as a side "
    "question that can be asked at any point. This use case suits a chatbot because reports are short, the information needed is structured and "
    "the routing rules are clear. The user journey has five steps. The user opens with a greeting or with a sentence describing the problem. "
    "The bot extracts what it can and asks only for what is missing. It then checks urgency, routes the report and confirms the outcome."
)
P(
    "Three principles shaped the wording. Each turn asks one question. Each reprompt tells the user what a good answer looks like. "
    "Every turn leaves a way to reach a person. These choices follow the techniques covered in class, which are intents with annotated "
    "training phrases, entities with synonyms, form-filling pages, conditional routes and event handlers for fallback."
)
H(3, "Implementation")
P(
    "The agent has one flow with a start page, a form page and five outcome pages (Figure {fig:flow}). It uses five custom intents, "
    "three custom entity types and a form with three required parameters. The two built-in default intents remain in "
    "place. Table {tab:map} maps each design element to the brief."
)
FIG("Dialogflow CX flow graph for the ESG Sustainability Support Assistant", Q1S / "01_flow_graph.png", 15.0,
    "Start Page routes to Collect Incident Details, Urgent Escalation, Human Handoff and Report Status Check. "
    "Collect Incident Details routes to Urgent Escalation, Accessibility Referral or Standard Confirmation. All outcome pages end the session.",
    key="flow")
T(
    "Mapping of the implementation to the brief",
    ["Requirement", "Implementation"],
    [
        ["3 to 5 intents", "report.incident (16 training phrases), report.urgent_hazard (14), request.human_agent (14), check.report_status (9) and ask.privacy (10). Default Welcome and Default Negative intents are also present."],
        ["Entity or parameter", "Entity types incident_type (water leak, energy waste, waste contamination, unsafe disposal, accessibility barrier, each with synonyms), facility_building (Buildings A to D, car park, main campus) and urgency_level (immediate, routine). The training phrases are annotated so a single sentence can fill the incident type and the building."],
        ["Branching", "On the Collect Incident Details page, a completed form goes to Urgent Escalation when urgency is immediate, to Accessibility Referral when the incident type is accessibility barrier, and otherwise to Standard Confirmation. The report.urgent_hazard intent skips the form and goes straight to Urgent Escalation."],
        ["Fallback", "No-match handlers on the start page and on each form parameter give a first and a second reprompt. The third no-match goes to Human Handoff. No-input has its own message."],
        ["Escalation", "Urgent Escalation and Human Handoff send a live agent handoff message after a short explanation. The prototype has no contact centre connected, so the handoff is a signal that a live deployment would pick up."],
    ],
    [3.2, 12.7],
    key="map",
)
P(
    "The form page asks for any parameter that is still missing. If the user opens with a sentence such as ‘There is a water "
    "leak in Building C that has been running all morning’, the incident type and building are filled from the sentence and "
    "only the urgency question remains (Figure {fig:prefill}). Figures {fig:urgent} to {fig:fallback} show the other paths, and Figure {fig:page} shows the form page. Table {tab:test} summarises the tests."
)
FIG("Prefilled parameters, form reprompt and routine confirmation", Q1S / "04_prefill_reprompt_routine.jpg", 15.0,
    "The simulator panel shows building and incident_type filled from the first message. The nonsense reply ‘banana pancakes’ triggers the form reprompt.",
    key="prefill")
FIG("Urgent path with the parameter table", Q1S / "03_urgent_path_parameters.png", 9.0,
    "The answer ‘immediate, water is flooding the corridor’ sends the user to Urgent Escalation and a person.", key="urgent")
FIG("Accessibility path", Q1S / "06_accessibility_path.png", 9.0,
    "An accessible entrance blocked for two days is classified as an accessibility barrier and referred to the Accessibility and Inclusion Office.", key="access")
FIG("Privacy question, repeated no-match and handoff", Q1S / "05_fallback_privacy_handoff.png", 9.0,
    "The privacy answer is given without leaving the conversation. Three consecutive unrecognised inputs lead to Human Handoff.", key="fallback")
FIG("Collect Incident Details page with its form parameters and routes", Q1S / "07_collect_incident_details_page.jpg", 15.0, key="page")
T(
    "Test log from the Dialogflow simulator",
    ["Test", "User inputs", "Expected result", "Observed result"],
    [
        ["Routine, prefilled", "Water leak in Building C message, then ‘banana pancakes’, then ‘it can wait, nobody is at risk’", "Type and building prefilled, reprompt on nonsense, routine confirmation", "As expected (Figure {fig:prefill})"],
        ["Routine, asked in turn", "Air conditioning message, ‘energy waste’, ‘Building B’, ‘routine’", "Three questions, then routine confirmation naming the issue and building", "As expected. Confirmation names energy waste and Building B."],
        ["Urgent", "Water leak message, then ‘immediate, water is flooding the corridor’", "Urgent Escalation with handoff and call 000 prompt", "As expected (Figure {fig:urgent})"],
        ["Accessibility", "Blocked accessible entrance message, ‘accessibility barrier’, ‘routine’", "Accessibility Referral", "As expected (Figure {fig:access})"],
        ["Privacy", "‘what data do you collect about me’", "Privacy answer, conversation stays open", "As expected (Figure {fig:fallback})"],
        ["Fallback to person", "Three nonsense inputs in a row", "Two reprompts, then Human Handoff", "As expected (Figure {fig:fallback})"],
    ],
    [2.6, 5.0, 4.6, 3.7],
    note="The urgent, accessibility and fallback screenshots were captured before the wording of the routine confirmation was shortened. That change touches only the Standard Confirmation page.",
    key="test",
)
H(3, "Conceptual link to a cloud workflow")
P(
    "The prototype stops at the chat. In a live deployment a webhook would pass each completed report to a message queue, and the queue would "
    "create a ticket, notify the right team and write a structured record that feeds an ESG dashboard. Urgent reports and requests for a person would "
    "trigger an on-call alert and a live agent handoff. Figure {fig:workflow} shows this design with Google Cloud and Azure service names. None of it "
    "has been built, and it is shown only to explain how the conversation data would become operational and reporting data."
)
FIG("Conceptual cloud workflow around the chatbot", ROOT / "q1_dialogflow" / "workflow_concept.png", 11.0, key="workflow")

H(2, "Q1(d) Accessibility and inclusion")
P(
    "The bot asks one short question per turn and uses plain words. It gives examples of acceptable answers, so a user does not need to "
    "know the categories in advance. The entity types carry synonyms such as ‘leaking’, ‘dripping tap’ and "
    "‘lights left on’, and the training phrases include short and informal wording, which helps users who write briefly or "
    "in non-standard English. A separate accessibility route sends barriers to the specialist office, and the person can ask for a human at any "
    "point. Because the conversation is text, it works with screen readers when it is placed in an accessible chat widget."
)
P(
    "These measures are design choices and have not been tested with assistive technology or with users who have disability. A live "
    "release would test the chat widget against the Web Content Accessibility Guidelines 2.2 (World Wide Web Consortium [W3C], 2023). "
    "It would also offer a non-chat route, such as a phone number or a form, for people who find chat difficult."
)

H(2, "Q1(e) Demo video")
P(
    "The video is aimed at a group of senior executives, so it opens with the business problem and closes with the next step. It is recorded in the "
    "Dialogflow simulator and runs for less than two minutes. The script in Table {tab:script} uses the same inputs as the tests above and takes about "
    "one minute and fifty seconds at a steady pace."
)
T(
    "Demo video script",
    ["Time", "On screen", "Narration"],
    [
        ["0:00 to 0:15", "Flow graph (Figure {fig:flow})", "Staff report leaks, wasted energy and access barriers in many ways, and reports often arrive late or incomplete. This assistant captures them in under a minute and sends them to the right team. It is built in Dialogflow CX."],
        ["0:15 to 0:50", "Simulator. Type the water leak message, then ‘banana pancakes’, then ‘it can wait, nobody is at risk’", "One sentence fills in the issue and the building, so the bot asks only for urgency. A nonsense reply gets a gentle reprompt. A routine answer ends with a confirmation, and the report is structured data that an ESG dashboard can use."],
        ["0:50 to 1:15", "Reset. Type the water leak message, then ‘immediate, water is flooding the corridor’", "An urgent answer takes a different branch. The bot tells the user to call 000 and connects them to a person, because safety should not wait for a form."],
        ["1:15 to 1:35", "Reset. Type ‘what data do you collect about me’", "The bot explains what it collects and asks users not to share personal details. Privacy and oversight are part of the design."],
        ["1:35 to 1:50", "Type three nonsense inputs, then show Figure {fig:workflow}", "If the bot cannot understand after two tries, it hands over to a person. The next step is a pilot at one site, with the bot connected to ticketing and to the ESG dashboard."],
    ],
    [2.4, 6.0, 7.5],
    key="script",
)
P(
    "The recording and the upload to Stream are still to be completed, and the link belongs on the title page.",
    indent=True,
)

page_break()
# --------------------------------------------------------------------------------------
# Q2
# --------------------------------------------------------------------------------------
H(1, "Q2 Cloud Security Incident Analysis")
H(2, "Q2(a) Incident summary")
P(
    "The incident is the 2024 theft of customer data from accounts on the Snowflake cloud data platform. Mandiant (2024) tracked the "
    "attacker as UNC5537 and reported that about 165 organisations were potentially exposed. The attacker logged in to customer "
    "Snowflake accounts with valid usernames and passwords that had been stolen earlier by infostealer malware. Many of the infected "
    "devices belonged to contractors who also used them for personal activity such as gaming and downloading pirated software. Some "
    "infections dated back to 2020, and at least 79.7 per cent of the accounts used had credentials that had been exposed before."
)
P(
    "The affected accounts had no multi-factor authentication. They did not limit sign-in to trusted network locations, and some "
    "credentials had not been rotated for as long as four years (Mandiant, 2024). The attacker ran queries, copied the results into "
    "temporary stages, downloaded the files and used the data to extort the victims. One victim, AT&T, reported that call and text "
    "records for about 110 million people were taken. The records covered May to October 2022 and one day in January 2023. They held "
    "phone numbers, interaction counts and call duration, but not message content or customer names (Security Affairs, 2024)."
)
T(
    "Incident summary",
    ["Question", "Answer"],
    [
        ["What happened", "An attacker used stolen credentials to log in to customer Snowflake accounts, copied large tables and extorted the owners. The activity became public in mid-2024."],
        ["Who was affected", "About 165 organisations were potentially exposed (Mandiant, 2024). AT&T reported call and text records for about 110 million people (Security Affairs, 2024)."],
        ["What was compromised", "Customer data held in tables in the victims’ Snowflake accounts. In the AT&T case this was call and text metadata, not message content."],
        ["Root cause", "Weak customer-side identity controls. Passwords alone were enough to sign in, because MFA was not enforced and no network allow-list was in place. Credentials stolen from infected contractor devices were not rotated. Mandiant found no evidence of a breach of Snowflake’s own enterprise environment."],
    ],
    [3.2, 12.7],
    key="incident",
)

H(2, "Q2(b) Cloud components, service model and shared responsibility")
P(
    "Figure {fig:q2arch} shows the components and the path of the attack. Snowflake is consumed as a managed data platform in a public cloud, where many "
    "customers share the same service and each customer holds its own accounts. The customer does not "
    "install or patch the database engine, so this report treats it as software as a service (SaaS). The platform runs on infrastructure "
    "as a service (IaaS) from AWS, Microsoft Azure or Google Cloud, which the customer never sees. Some writers call a data platform of this "
    "kind platform as a service (PaaS) because customers load and query their own data. The classification does not change the analysis "
    "of responsibility, because in both models the customer keeps control of identities, data and the devices that connect."
)
FIG("Cloud components and attack path in the Snowflake incident", ROOT / "q2_snowflake" / "snowflake_architecture.png", 16.0,
    "Shaded areas show who is responsible for each component. Numbers show the order of the attack. Based on Mandiant (2024).", key="q2arch")
P(
    "Under the shared responsibility model, the provider secures the service and the infrastructure beneath it, and the customer secures what "
    "it puts into the service and how people access it (Microsoft, n.d.-a). In a SaaS service the customer always keeps responsibility for "
    "its data, its accounts and identities, and the devices that connect. Table {tab:shared} applies the model to this incident."
)
T(
    "Shared responsibility applied to the incident",
    ["Layer", "Who is responsible", "Status in the incident"],
    [
        ["Physical data centres, hosts and network", "Cloud provider (IaaS vendor and Snowflake)", "No evidence of failure (Mandiant, 2024)"],
        ["Snowflake service, query engine and authentication service", "Snowflake", "No evidence of breach of Snowflake’s enterprise environment (Mandiant, 2024)"],
        ["User accounts, passwords and MFA settings", "Customer", "Failed. Password-only sign-in and unrotated credentials."],
        ["Network access policy (allowed addresses)", "Customer", "Failed. No allow-list in the affected accounts."],
        ["Contractor and staff devices", "Customer", "Failed. Infostealers ran on devices used for personal activity."],
        ["Data classification and monitoring of data access", "Customer", "No evidence that large exports were flagged while they ran."],
    ],
    [5.2, 4.6, 6.1],
    key="shared",
)
P(
    "Snowflake has since changed its defaults. Its documentation sets out a phased move that makes MFA mandatory for people who sign in with "
    "a password and blocks password-only sign-in for service users (Snowflake, n.d.). The change shows that the provider can reduce customer "
    "risk, but the customer still has to configure and operate the controls."
)

H(2, "Q2(c) Prevention steps")
P(
    "Table {tab:prevent} lists five steps and names the platform tools that carry them out. Each step breaks at least one numbered stage in Figure {fig:q2arch}."
)
T(
    "Prevention steps and the tools that implement them",
    ["Step", "Tools", "Attack stage it breaks"],
    [
        ["1 Require MFA for every human user and use single sign-on", "Snowflake authentication policies that require MFA, or federated sign-in through Microsoft Entra ID with a Conditional Access policy that demands MFA", "Stage 4. A stolen password alone no longer opens the account."],
        ["2 Restrict where sign-in can occur", "Snowflake network policies with an IP allow-list. Private connectivity through AWS PrivateLink or Azure Private Link for internal users", "Stages 4 and 6. A valid login from an unknown VPN address is refused."],
        ["3 Remove long-lived passwords and rotate secrets", "Key-pair or OAuth authentication for service users. AWS Secrets Manager or Azure Key Vault with scheduled rotation. Monitoring of exposed credentials", "Stages 3 and 4. Stolen credentials go stale or never work for service accounts."],
        ["4 Detect abnormal access and respond fast", "Snowflake LOGIN_HISTORY and ACCESS_HISTORY views, exported to Microsoft Sentinel, with alert rules for new locations and large COPY INTO or GET activity. A response playbook that disables the user and rotates credentials", "Stages 7 and 8. Exfiltration is flagged within minutes and the account can be disabled."],
        ["5 Protect the devices that hold credentials", "Microsoft Defender for Endpoint and Intune compliance policies, with Conditional Access that blocks non-compliant devices. Separate work and personal devices for contractors", "Stage 1. Infostealers are blocked or detected, and an unmanaged device cannot reach the account."],
    ],
    [4.2, 7.2, 4.5],
    key="prevent",
)
P(
    "Steps 1 and 2 address the two missing controls that Mandiant (2024) identified in the affected accounts. Steps 3 to 5 "
    "lower the chance and the damage of a repeat by shortening the life of stolen credentials and the time taken to detect misuse."
)
H(3, "Governance, incident response and data protection")
P(
    "The steps above sit inside a governance frame. A client should treat each contractor and each data platform as a third party with its own risk "
    "assessment, require managed devices for anyone who touches production data, and review account access on a fixed schedule. It should also classify "
    "sensitive columns and apply masking policies, such as Snowflake dynamic data masking, so that one compromised account sees less. For response, the "
    "client needs a written playbook that covers disabling accounts, rotating credentials, preserving query logs for investigation and deciding who "
    "speaks to regulators and customers. An Australian organisation that suffers an eligible data breach must also assess it and notify the Office of the "
    "Australian Information Commissioner and affected individuals under the Notifiable Data Breaches scheme (Office of the Australian Information Commissioner [OAIC], n.d.)."
)

page_break()
# --------------------------------------------------------------------------------------
# Q3
# --------------------------------------------------------------------------------------
H(1, "Q3 LLM Triage of ESG Messages")

runs = [json.loads(l) for l in open(ROOT / "q3_llm_triage" / "results" / "llm_runs.jsonl")]


def get(ver, mid):
    return next(r for r in runs if r["prompt_version"] == ver and r["id"] == mid)


H(2, "Q3(a) Prompt critique, revised prompt and sample runs")
P(
    "The original prompt asks for eight fields and tells the model to return JSON only. It has six weaknesses. It does not list the "
    "allowed values for issue_category, recommended_team or data_sensitivity_risk, so the model invents labels. It gives no urgency "
    "rubric, so CRITICAL and HIGH depend on the model’s own judgement. It does not say what to do when information is missing, which "
    "invites guessing. It has no instruction to ignore text inside the message that looks like a command. It has no rule that keeps urgency "
    "independent of the reporter’s role, tone or grammar. Finally, it has no example, no schema and no field to flag cases for human review."
)
P(
    "The revised prompt is a system prompt with a role, closed value lists, an urgency rubric, six rules, a JSON schema and two worked examples. "
    "It adds human_review_required, location_mentioned and confidence, and it wraps the message in tags so the model treats it as data. "
    "The full text is in the repository (prompts folder)."
)
P(
    "Both prompts were run on Qwen2.5-1.5B-Instruct (Qwen Team, 2024) on a CPU, with greedy decoding and no sampling. A small open model "
    "was used because it could be run locally and logged in full. Results for a larger commercial model may differ. Ten messages were "
    "used. Five come from the assessment brief and five were written for this experiment to cover praise, conduct, safety and a "
    "question. The reference labels were assigned by the analyst, not by a domain expert."
)
P("Three messages are shown here, and all 28 runs (20 main and 8 bias probes) are in the repository. The first is a facilities problem, the second is a conduct concern and the third is a gas smell.")
for mid in ("m1", "m7", "m8"):
    r1 = get("v1", mid)
    P(f"**Message {mid}.** {r1['text']}", indent=False, after=2, keep=True)
    CODE(json.dumps(r1["parsed"], indent=1))
P("For comparison, the original prompt produced the following for the gas message (m8). The output is wrapped in markdown fences, uses free-text labels and omits human_review_required.", indent=True)
CODE(get("v0", "m8")["raw_output"].strip())
P(
    "The revised prompt fixed the format. All ten outputs parsed, used only allowed values and contained no fences (Table {tab:format}). "
    "The original prompt gave valid enum values in none of ten runs, and it rated the gas smell as HIGH instead of CRITICAL."
)
P(
    "Table {tab:acc} compares accuracy across all four systems. The revised prompt did not remove every error. The model sent the recycling bin report to the Accessibility and Inclusion Office and the air conditioning report to Procurement. It labelled the "
    "conduct concern (m7) as health and safety and also sent it to the Accessibility and Inclusion Office. It marked data sensitivity as HIGH for five "
    "of ten messages, including the gas smell, which names no person. Its summaries contain invented words, such as ‘urgent evacuation "
    "required’ for the gas message and ‘Customer’ for an employee question. Its confidence was 0.95 on the supplier and conduct messages, "
    "which it categorised wrongly, so the confidence value is not a usable signal for this model."
)

H(2, "Q3(b) Comparison with baselines")
P(
    "Two baselines were run on the same ten messages and the eight bias probes. The rule-based baseline uses ordered keyword rules, with safety words first, "
    "and maps each category to a team. The zero-shot baseline uses facebook/bart-large-mnli (Lewis et al., 2020), which scores a message against "
    "a natural-language description of each category and each urgency level. The rule set was written by the analyst with the sample "
    "messages in view, so its scores are optimistic. It would probably fall on unseen wording."
)
a = json.load(open(ROOT / "q3_llm_triage" / "results" / "analysis.json"))
T(
    "Output format discipline on ten messages",
    ["Prompt", "Valid allowed values", "Wrapped in markdown fences", "Missing human_review_required"],
    [
        ["Original (v0)", "0 of 10", "10 of 10", "10 of 10"],
        ["Revised (v1)", "10 of 10", "0 of 10", "0 of 10"],
    ],
    [4, 4, 4, 4],
    key="format",
)
acc = a["accuracy"]
rows = []
for key_, label in (("llm_v0", "LLM original prompt"), ("llm_v1", "LLM revised prompt"), ("rule_based", "Rule-based keywords"), ("zeroshot_bart_mnli", "Zero-shot bart-large-mnli")):
    d = acc[key_]
    rows.append([label, f"{d['issue_category']}/10", f"{d['urgency']}/10", f"{d['followup_required']}/10", f"{d['recommended_team']}/10", f"{d['urgency_within_one']}/10"])
T(
    "Agreement with the analyst reference labels on ten messages",
    ["System", "Category", "Urgency", "Follow-up", "Team", "Urgency within one level"],
    rows,
    [4.6, 2.0, 2.0, 2.2, 2.0, 3.1],
    note="The original prompt scores 0 of 10 on category and team because it produced free-text labels outside the allowed list, which says little about its understanding of the message.",
    key="acc",
)
P(
    "On urgency the revised LLM matched the reference on eight of ten messages, the zero-shot baseline on six, and the rule baseline on all ten. "
    "The LLM and the zero-shot baseline both rated the gas smell CRITICAL, but the zero-shot model rated the blocked accessible entrance LOW, which "
    "would delay a barrier to access. On category, the LLM and the baselines each matched seven of ten. The LLM mislabelled the supplier concern as "
    "governance and the conduct concern as health and safety. The rule baseline missed the supplier concern and the composting praise because its keywords did not "
    "match the word forms ‘suppliers’ and ‘composting’, and it also labelled the conduct concern as health and safety. The zero-shot model labelled the recycling bin, air conditioning and coffee cup messages as general enquiries."
)
P(
    "On consistency, the greedy LLM runs were not repeated, so repeat-run stability is untested. The consistency measure used here is how stable each system is "
    "across rewordings of the same facts (Table {tab:bias}). On the water leak, the revised LLM gave HIGH for the neutral and junior variants and CRITICAL for "
    "the senior reporter and the non-standard English variants. The rule-based baseline gave MEDIUM for three variants and HIGH for one. The zero-shot baseline "
    "gave HIGH for all four. On the blocked ramp, the LLM and the rule baseline gave HIGH for all four, and the zero-shot baseline switched between LOW and HIGH."
)
bp = a["bias"]
rows = []
for k, v in bp.items():
    rows.append([k.replace("_", " "), v["variant"], v["llm"], v["rule"], v["zs"]])
T(
    "Urgency by reporter variant for the same facts",
    ["Probe", "Variant", "LLM revised", "Rule-based", "Zero-shot"],
    rows,
    [3.0, 4.2, 2.8, 3.0, 2.9],
    note="The b1 probe describes a ceiling water leak. The b2 probe describes a blocked accessible ramp. Each variant was run once.",
    key="bias",
)
P(
    "The probe shows a pattern worth acting on, but it does not prove bias. The sample is eight messages and one run each. The LLM raised "
    "urgency for a senior job title and for terse, non-standard wording even though the prompt told it to ignore both. The rule baseline did "
    "not react to the job title, but it reacted to word forms in the same way. The zero-shot model reacted to added words. Each system has a different "
    "weakness. The rule baseline is predictable but brittle, the zero-shot model is cheap and unstable on urgency, and the LLM follows the format "
    "well but can move urgency on surface features. A mix of the three, with a rule layer for safety words, is the more defensible design."
)
P(
    "On cost and speed, the rule baseline runs in under a millisecond. The zero-shot model took about 3 seconds per message and the "
    "revised LLM about " + str(round(a["latency_v1"])) + " seconds per message on the CPU used, which a hosted model would cut sharply."
)

H(2, "Q3(c) Improvements and production readiness")
P(
    "The pipeline is not ready for production as an automatic router. The evidence is ten messages, one small model and analyst-made labels. A "
    "staged release would use it only as decision support, with a person reviewing every CRITICAL and HIGH-sensitivity case. Before wider use the team should "
    "build a labelled set of several hundred real messages with two reviewers, report recall for CRITICAL cases separately, and set a release "
    "threshold. It should add a safety-word override so that terms such as gas, fire and injury always reach the emergency team, and it should "
    "use schema-enforced output so malformed JSON cannot occur."
)
P(
    "Hallucination and data protection need their own controls. The invented evacuation advice in the gas summary shows the hallucination risk, so free-text "
    "summaries should be labelled as machine-generated and never trigger action alone. The model here ran locally, so no message left the machine. A hosted API would send "
    "message text, which can name staff, to an external service, so the team would need personal data redaction, a contract that limits retention and "
    "training use, and a data region that meets policy. The team should repeat the paired-variant bias test on every prompt or model change, "
    "test for prompt injection, and monitor drift and reviewer overrides after launch."
)

H(2, "Q3(d) Azure build-out")
P(
    "Figure {fig:azure} shows a design for Azure. Staff submit messages through Microsoft Forms, a Teams channel or a shared mailbox. A Logic App picks up each message "
    "and puts it on an Azure Service Bus queue, which absorbs bursts and holds failed items in a dead-letter queue. An Azure Function takes each item and calls "
    "Azure AI Language to redact personal data (Microsoft, n.d.-d). It then calls Azure AI Content Safety Prompt Shields to check for injected instructions "
    "(Microsoft, n.d.-c) before it calls an Azure OpenAI deployment that returns the triage JSON under a schema-enforced structured output setting (Microsoft, n.d.-b)."
)
FIG("Proposed Azure architecture for ESG message triage", ROOT / "q3_llm_triage" / "diagrams" / "azure_architecture.png", 11.5, key="azure")
P(
    "A validation step in the Function checks the JSON against the schema and applies the safety-word override and the confidence threshold. Valid, "
    "confident items go to the service desk or Planner. CRITICAL items, HIGH-sensitivity items and any item that fails validation go to a human review "
    "queue in Teams as an adaptive card. Every input, redacted text, output and reviewer decision is stored in Azure Cosmos DB, and Power BI reads from "
    "it to show volumes, response times and override rates. Managed identities and Azure Key Vault remove stored secrets, private endpoints keep traffic "
    "off the public internet, Azure Monitor and Application Insights record failures and latency, and Microsoft Purview and Microsoft Entra ID role-based "
    "access control govern who can read the data."
)

page_break()
# --------------------------------------------------------------------------------------
# Q4
# --------------------------------------------------------------------------------------
H(1, "Q4 NotebookLM Evaluation")
P(
    "This section evaluates NotebookLM as a tool for a university pilot. NotebookLM is Google’s source-grounded research assistant. In the session used "
    "for this report the interface was labelled Gemini Notebook, at notebook.google.com. A notebook holds a set of sources, and the assistant answers from those sources."
)
H(2, "Q4(a) Features")
P(
    "Table {tab:features} lists the features that suit common academic activities and the university users who would use each one."
)
T(
    "NotebookLM features and their academic uses",
    ["Feature", "What it does", "Academic use", "Typical user"],
    [
        ["Source upload and source guide", "Adds web pages, files or text as sources and shows the source text beside the chat.", "Collecting readings for one topic in one place.", "Student, researcher"],
        ["Chat with citations", "Answers questions from the sources with numbered citations. Selecting a citation opens the supporting passage.", "Finding and checking facts in long readings.", "Student, researcher, lecturer"],
        ["Briefing document (Reports)", "Generates a structured brief with an executive summary and key takeaways.", "Writing a research summary or a workshop briefing.", "Researcher, lecturer"],
        ["Mind map", "Turns the sources into an expandable tree of topics.", "Seeing the structure of a topic and extracting themes.", "Student, researcher"],
        ["Flashcards", "Generates question and answer cards.", "Revising definitions and key facts for an exam.", "Student"],
        ["Quiz", "Generates multiple-choice questions with feedback on every option.", "Self-testing, or drafting practice questions for a class.", "Student, lecturer"],
        ["Audio Overview", "Generates a spoken summary in a chosen format and length.", "Reviewing a topic while travelling.", "Student"],
    ],
    [3.6, 5.6, 4.1, 2.6],
    note="The Studio panel also offers slide decks, video overviews, infographics and data tables. These were not tested.",
    key="features",
)

H(2, "Q4(b) Demonstration in a scenario")
P(
    "The scenario is a postgraduate student who must prepare a briefing for a workshop on ESG reporting standards. The student collects three public pages from the IFRS Foundation. "
    "These are the introduction to the ISSB and IFRS Sustainability Disclosure Standards, IFRS S1 General Requirements for Disclosure of Sustainability-related Financial Information, and IFRS S2 Climate-related "
    "Disclosures (IFRS Foundation, n.d.-a, n.d.-b, n.d.-c). The same notebook shows how a lecturer could draw on the quiz and the flashcards for class practice, and how a "
    "researcher could use the mind map to extract themes. Attempts to add pages from Australian regulators failed because those sites could not be fetched from "
    "the environment used, so the notebook has no Australian source. That gap affects the evaluation and is discussed in Q4(c)."
)
H(3, "Chat with citations")
P(
    "The first prompt was ‘What are the four core content areas that IFRS S1 and S2 require entities to disclose, and when do the standards take effect?’ "
    "The answer named governance, strategy, risk management, and metrics and targets, and gave 1 January 2024 as the start of the annual reporting "
    "periods in which the standards apply, with earlier application permitted if S2 is applied at the same time (Figure {fig:q4cited}). The second prompt asked about Scope 3 "
    "emissions. The answer said S2 requires Scope 1, 2 and 3 emissions to be measured under the Greenhouse Gas Protocol, and that a first-year relief applies "
    "(Figure {fig:q4scope3}). Selecting a citation opened the exact passage in the source (Figures {fig:q4cite} and {fig:q4trans}), which makes each claim checkable in seconds."
)
FIG("Chat answer on the four content areas and effective date, with citations", Q4S / "02_cited_answer_four_pillars.jpg", 14.5, key="q4cited")
FIG("Chat answer on Scope 3 emissions disclosure", Q4S / "03_scope3_answer.jpg", 14.5, key="q4scope3")
FIG("A citation opens the supporting passage in the source", Q4S / "04_citation_opens_source_passage.jpg", 14.5, key="q4cite")
FIG("A citation preview for the first-year transitional relief", Q4S / "05_transitional_relief_citation.jpg", 14.5, key="q4trans")
H(3, "Stress tests for hallucination and scope")
P(
    "Three further prompts tested whether the tool stays inside its sources. A question about ASIC penalties for failing to comply with mandatory climate reporting returned a refusal. "
    "The assistant said the sources do not cover ASIC or penalties, explained what they do cover, and offered to run web research (Figure {fig:q4asic}). "
    "A second prompt asked which paragraph of IFRS S2 sets a mandatory carbon price of 50 dollars per tonne. The premise is false, and the assistant said so. It explained "
    "that S2 asks entities to disclose the internal carbon prices they use and listed the cross-industry metric categories (Figure {fig:q4false}). A third prompt asked how many Scope 3 categories "
    "the GHG Protocol defines and which one covers employee commuting. The assistant gave the number, 15, which the source states, and declined to name the commuting category because "
    "the sources do not list the categories (Figure {fig:q4cat}). The tool therefore declined to answer or corrected the premise instead of inventing an answer."
)
FIG("Refusal of a question the sources do not cover", Q4S / "06_out_of_scope_ASIC_refusal.jpg", 14.5, key="q4asic")
FIG("Rejection of a false premise about a carbon price", Q4S / "14_false_premise_carbon_price_rejected.jpg", 14.5, key="q4false")
FIG("Answer that stays inside the sources on Scope 3 categories", Q4S / "15_scope3_categories_declines_commuting.jpg", 14.5, key="q4cat")
H(3, "Briefing document and mind map")
P(
    "The briefing document opened with an executive summary and key takeaways that covered the four-pillar structure, scope and materiality, climate metrics under the GHG Protocol and "
    "transitional relief (Figure {fig:q4brief}). The mind map placed the four pillars as branches under a root node on the two standards (Figure {fig:q4mind}). Each pillar expands into two or three child nodes, such as "
    "Assessment Processes and GHG Protocol Integration."
)
FIG("Briefing document generated from the three sources", Q4S / "07_briefing_document.jpg", 14.5, key="q4brief")
FIG("Mind map with the four pillars expanded", Q4S / "08_mind_map.jpg", 14.5, key="q4mind")
H(3, "Flashcards, quiz and audio overview")
P(
    "The flashcard deck had 60 cards. The first card asked for the two standard-setting boards housed in the IFRS Foundation, and its answer, the "
    "International Accounting Standards Board and the International Sustainability Standards Board, is correct (Figures {fig:q4fcq} and {fig:q4fca}). The quiz had ten questions. "
    "The first asked what condition applies to early adoption of IFRS S1. The correct option, that S2 must be applied at the same time, was marked right, and each wrong "
    "option came with a reason (Figures {fig:q4quizq} and {fig:q4quizf}). An Audio Overview in the Brief format was generated, and it played back for 1 minute 47 seconds (Figure {fig:q4studio}). Its spoken content was not "
    "audited line by line for this report, so no claim is made about its accuracy."
)
FIG("Flashcard question", Q4S / "09_flashcard_question.jpg", 12.5, key="q4fcq")
FIG("Flashcard answer", Q4S / "10_flashcard_answer.jpg", 12.5, key="q4fca")
FIG("Quiz question", Q4S / "11_quiz_question.jpg", 12.5, key="q4quizq")
FIG("Quiz feedback on every option", Q4S / "12_quiz_feedback.jpg", 12.5, key="q4quizf")
FIG("Studio panel with the generated audio overview", Q4S / "13_audio_overview_ready_studio_list.jpg", 14.5, key="q4studio")

H(2, "Q4(c) Critical analysis")
P(
    "Table {tab:percap} assesses each capability on accuracy and relevance, usefulness in academic work, and limitations or concerns. "
    "The evidence is the set of checks in Table {tab:checks}, which compares ten outputs with the IFRS pages. Nine were accurate. The executive summary of the briefing "
    "document said the standards ‘mandate’ reporting, which is imprecise, because the ISSB sets the requirements and each jurisdiction decides whether to make them compulsory. "
    "The sample is small and covers short factual claims. It does not show how the tool performs on long or ambiguous passages."
)
T(
    "Assessment of each capability",
    ["Capability", "Accuracy and relevance", "Usefulness in academic work", "Limitations and concerns"],
    [
        ["Source upload", "The source text was imported in full and is shown beside the chat, so passages can be checked.", "Keeps a reading set in one place and makes checking quick.", "The web import carried the site’s menus and links into the source, which adds noise. Only public pages were used."],
        ["Chat with citations", "Four factual checks were accurate. A false premise was corrected. Citations opened the supporting passage.", "Saves time on long readings and supports referencing because each claim points to a passage.", "Answers can be cautious. The tool would not name the commuting category although the topic is standard. Quality depends on the sources, and there was no Australian source."],
        ["Briefing document", "Content matched the sources, with one imprecise word (‘mandate’).", "Gives a usable first draft of a summary for a workshop or a literature note.", "Loose phrasing can pass into a student’s own writing. The long document was not checked line by line."],
        ["Mind map", "The four pillars matched the sources.", "Shows the structure of a topic and helps with theme extraction.", "Shallow, and some node labels repeat under different pillars, so it adds little beyond the headings."],
        ["Flashcards", "Card 1 was accurate. Only one of 60 cards was checked.", "Fast revision material in the exact language of the sources.", "Fifty-nine cards were not checked, so a student could memorise an error. Risk of over-reliance."],
        ["Quiz", "Question 1 was accurate and each option carried an explanation.", "Useful for self-testing, and a lecturer could adapt the questions.", "Only one of ten questions was checked. Distractors may be weak or the keyed answer may be wrong."],
        ["Audio Overview", "Generated and played back. Content not audited.", "Suits revision while travelling.", "Spoken claims cannot be checked against sources line by line as a cited answer can."],
    ],
    [2.6, 4.4, 4.2, 4.7],
    size=9,
    key="percap",
)
T(
    "Spot checks of NotebookLM outputs against the sources",
    ["Output", "Claim checked", "Result"],
    [
        ["Chat answer", "Four content areas are governance, strategy, risk management, and metrics and targets", "Accurate"],
        ["Chat answer", "Standards apply to annual periods beginning on or after 1 January 2024, earlier application permitted", "Accurate"],
        ["Chat answer", "S2 requires Scope 1, 2 and 3 emissions measured under the GHG Protocol", "Accurate"],
        ["Chat answer", "First-year transitional relief for Scope 3 in the first annual reporting period", "Accurate"],
        ["Chat answer", "S2 sets no fixed carbon price and asks for internal carbon prices used", "Accurate"],
        ["Chat answer", "GHG Protocol defines 15 Scope 3 categories", "Accurate"],
        ["Flashcard 1", "IASB and ISSB sit within the IFRS Foundation", "Accurate"],
        ["Quiz question 1", "Early application of S1 requires S2 to be applied at the same time", "Accurate"],
        ["Mind map", "Four pillars of governance, strategy, risk management, and metrics and targets", "Accurate"],
        ["Briefing document", "Executive summary says the standards mandate reporting", "Imprecise wording"],
    ],
    [3.0, 9.5, 3.4],
    key="checks",
)
P(
    "**Adoption for the university pilot.** The evidence supports a limited pilot. The tool saves time, its citations make checking possible, and it declined to answer instead "
    "of inventing a fact. A pilot should set conditions. Students and staff should check any figure against the cited passage before they use it. Courses should say how "
    "generated summaries may be used in assessed work, because a polished summary can hide weak understanding. The university should confirm Google’s data terms for the account type "
    "used before any unpublished or personal material is uploaded, and it should check that the interface works with screen readers and keyboard navigation. Access to a "
    "free tool can change, so a course should not depend on one feature staying available."
)

page_break()
# --------------------------------------------------------------------------------------
# AI acknowledgement
# --------------------------------------------------------------------------------------
H(1, "AI Acknowledgement")
P(
    "The applicable category is AI Exploration. I used Claude, an AI model from Anthropic, as my assistant throughout this assessment. I set the "
    "direction and made the decisions. Claude assisted by carrying out the Dialogflow build steps, running the Q3 and Q4 experiments, "
    "producing the diagrams and drafting text, all under my instructions. I reviewed the work and take responsibility for what is submitted. The prompts, "
    "scripts, outputs and screenshots are in the GitHub repository.",
    indent=False,
)

# --------------------------------------------------------------------------------------
# References
# --------------------------------------------------------------------------------------
page_break()
H(1, "References")
refs = [
    "IFRS Foundation. (n.d.-a). *IFRS S1 general requirements for disclosure of sustainability-related financial information*. https://www.ifrs.org/issued-standards/ifrs-sustainability-standards-navigator/ifrs-s1-general-requirements/",
    "IFRS Foundation. (n.d.-b). *IFRS S2 climate-related disclosures*. https://www.ifrs.org/issued-standards/ifrs-sustainability-standards-navigator/ifrs-s2-climate-related-disclosures/",
    "IFRS Foundation. (n.d.-c). *Introduction to the ISSB and IFRS sustainability disclosure standards*. https://www.ifrs.org/sustainability/knowledge-hub/introduction-to-issb-and-ifrs-sustainability-disclosure-standards/",
    "Lewis, M., Liu, Y., Goyal, N., Ghazvininejad, M., Mohamed, A., Levy, O., Stoyanov, V., & Zettlemoyer, L. (2020). BART: Denoising sequence-to-sequence pre-training for natural language generation, translation, and comprehension. *Proceedings of the 58th Annual Meeting of the Association for Computational Linguistics*, 7871–7880. https://doi.org/10.18653/v1/2020.acl-main.703",
    "Mandiant. (2024, June 10). *UNC5537 targets Snowflake customer instances for data theft and extortion*. Google Cloud Blog. https://cloud.google.com/blog/topics/threat-intelligence/unc5537-snowflake-data-theft-extortion",
    "Microsoft. (n.d.-a). *Shared responsibility in the cloud*. Microsoft Learn. https://learn.microsoft.com/en-us/azure/security/fundamentals/shared-responsibility",
    "Microsoft. (n.d.-b). *How to use structured outputs with Azure OpenAI*. Microsoft Learn. https://learn.microsoft.com/en-us/azure/ai-foundry/openai/how-to/structured-outputs",
    "Microsoft. (n.d.-c). *What is Azure AI Content Safety?* Microsoft Learn. https://learn.microsoft.com/en-us/azure/ai-services/content-safety/overview",
    "Microsoft. (n.d.-d). *What is the personally identifiable information (PII) detection feature in Azure Language?* Microsoft Learn. https://learn.microsoft.com/en-us/azure/ai-services/language-service/personally-identifiable-information/overview",
    "Office of the Australian Information Commissioner. (n.d.). *Part 4: Notifiable Data Breach (NDB) scheme*. https://www.oaic.gov.au/privacy/privacy-guidance-for-organisations-and-government-agencies/preventing-preparing-for-and-responding-to-data-breaches/data-breach-preparation-and-response/part-4-notifiable-data-breach-ndb-scheme",
    "Privacy Act 1988 (Cth). https://www.legislation.gov.au/C2004A03712/latest/text",
    "Qwen Team. (2024). *Qwen2.5 technical report*. arXiv. https://arxiv.org/abs/2412.15115",
    "Security Affairs. (2024, July 12). *New AT&T data breach exposed call logs of almost all customers*. https://securityaffairs.com/165658/data-breach/att-disclosed-a-new-data-breach.html",
    "Snowflake. (n.d.). *Multi-factor authentication rollout* [Documentation page]. Snowflake Documentation. https://docs.snowflake.com/en/user-guide/security-mfa-rollout",
    "World Wide Web Consortium. (2023). *Web content accessibility guidelines (WCAG) 2.2*. https://www.w3.org/TR/WCAG22/",
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-1.27)
    p.paragraph_format.space_after = Pt(6)
    _runs(p, ref)

doc.save(OUT)
json.dump(NEWREG, open(REGFILE, "w"))
print("saved", OUT)
